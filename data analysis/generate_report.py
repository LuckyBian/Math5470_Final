from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Title
title = doc.add_heading('Home Credit Default Risk Data Analysis Report', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1. Introduction
doc.add_heading('1. Introduction', level=1)
doc.add_paragraph(
    "This analysis aims to identify key drivers of default risk within the Home Credit dataset. "
    "We examined the target distribution, data quality, and relationships between various applicant features "
    "and their repayment behavior. The insights derived here serve as a foundation for feature engineering and model selection."
)

# 2. Target Variable Analysis
doc.add_heading('2. Target Variable Analysis', level=1)
doc.add_paragraph(
    "The dataset exhibits a significant class imbalance. The default rate (Target=1) is approximately 8.07%, "
    "while the vast majority of applicants (91.93%) successfully repaid their loans. This disparity suggests that "
    "metric selection should prioritize Area Under the Curve (AUC) rather than accuracy. Stratified K-Fold cross-validation "
    "is recommended to ensure stable validation sets that reflect this distribution."
)

# 3. Data Quality Assessment
doc.add_heading('3. Data Quality Assessment', level=1)
doc.add_paragraph(
    "A substantial portion of the dataset contains missing values, particularly in housing-related variables. "
    "Features such as COMMONAREA_AVG, NONLIVINGAPARTMENTS_AVG, and FONDKAPREMONT_MODE have missing rates exceeding 68%. "
    "These high missing rates imply that housing information is not mandatory for all applicants. Standard imputation strategies "
    "might be insufficient here, so treating 'missing' as a separate category or using tree-based models that handle missing values natively is advisable."
)

# 4. Demographic Analysis
doc.add_heading('4. Demographic Analysis', level=1)
doc.add_paragraph("Applicant demographics show clear correlations with default risk.")
p = doc.add_paragraph()
p.add_run("Age: ").bold = True
p.add_run(
    "There is a negative correlation between age and default probability. Younger applicants demonstrate a "
    "significantly higher likelihood of default compared to older applicants."
)
p = doc.add_paragraph()
p.add_run("Employment: ").bold = True
p.add_run(
    "A similar trend is observed with employment duration. Applicants with longer tenure in their current employment "
    "exhibit lower default rates."
)
p = doc.add_paragraph()
p.add_run("Gender: ").bold = True
p.add_run(
    "A distinct disparity exists between genders. Male applicants have a default rate of approximately 10.14%, "
    "whereas female applicants have a lower rate of roughly 7.00%."
)

# 5. Financial and External Score Analysis
doc.add_heading('5. Financial and External Score Analysis', level=1)
doc.add_paragraph(
    "External credit scores provided in the dataset (EXT_SOURCE_1, EXT_SOURCE_2, EXT_SOURCE_3) are among the strongest "
    "predictors of risk. All three show a robust negative correlation with the target (ranging from -0.15 to -0.18), "
    "indicating that higher external scores reliably predict lower default risk. Additionally, the 'Credit-to-Income Ratio' "
    "analysis reveals that applicants with higher debt leverage are marginally riskier, though extreme outliers in income "
    "can distort this signal if not capped."
)

# 6. Behavioral History
doc.add_heading('6. Behavioral History', level=1)
doc.add_paragraph("Historical data from previous applications and credit bureaus provides strong behavioral signals.")
p = doc.add_paragraph()
p.add_run("Previous Refusals: ").bold = True
p.add_run(
    "Applicants who have been refused loans in the past show a markedly higher probability of defaulting on current loans."
)
p = doc.add_paragraph()
p.add_run("Active Credits: ").bold = True
p.add_run(
    "A higher number of active credits in the Credit Bureau system correlates with increased risk, likely due to a higher total debt burden."
)
p = doc.add_paragraph()
p.add_run("Payment History: ").bold = True
p.add_run(
    "Analysis of installment payments confirms that past overdue behavior is highly predictive. Applicants with a history of "
    "late payments (high 'Days Past Due') are significantly more likely to default again."
)

# 7. Contract Type Analysis
doc.add_heading('7. Contract Type Analysis', level=1)
doc.add_paragraph(
    "The type of loan contract also influences risk. 'Cash loans' are associated with a higher default rate (8.35%) "
    "compared to 'Revolving loans' (5.48%). This difference suggests that the nature of the financial product itself "
    "acts as a risk filter."
)

# 8. Conclusion
doc.add_heading('8. Conclusion', level=1)
doc.add_paragraph(
    "The analysis confirms that EXT_SOURCE scores, age, and historical repayment behavior are the dominant predictors of default risk. "
    "The severe class imbalance and high missing data rates in housing features require careful preprocessing. Future modeling efforts "
    "should focus on interaction features between age and employment, as well as aggregating historical behavioral data to capture "
    "applicant creditworthiness fully."
)

# Save the document
output_path = '/data/weizhen/code/math/Home_Credit_Analysis_Report.docx'
doc.save(output_path)
print(f"Report saved to {output_path}")





